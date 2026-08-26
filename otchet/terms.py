"""Замена раздела «Термины и определения» в готовом отчёте.

Скрипт правит Отчет_заполненный.docx на месте, а не пересобирает его из
шаблона: в выходном файле есть ручные правки, которые пересборка потеряла бы.
Термины первого этапа (чат-бот, предиктивная аналитика, Python) заменены на
понятия, без которых не читаются разделы 1.3 и 1.5.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = Path(__file__).resolve().parent.parent / "Отчет_заполненный.docx"
FONT, SIZE = "Times New Roman", Pt(14)

TERMS = [
    ("CRM (от англ. Customer Relationship Management)",
     " – система управления взаимоотношениями с клиентами. Представляет собой "
     "программу, которая помогает бизнесу управлять всеми этапами взаимодействия "
     "с клиентами: от сбора контактов и ведения клиентской базы до управления "
     "продажами и анализа показателей эффективности."),
    ("Коллаборативная фильтрация",
     " – метод формирования рекомендаций, основанный на сходстве поведения "
     "пользователей: предпочтения одного восстанавливаются по выбору других "
     "пользователей со схожей историей. Не требует описания свойств "
     "рекомендуемых объектов и опирается только на историю взаимодействий."),
    ("Матричное разложение (SVD, от англ. Singular Value Decomposition)",
     " – способ реализации коллаборативной фильтрации, при котором разреженная "
     "матрица взаимодействий «пользователь – объект» представляется "
     "произведением матриц меньшей размерности. Полученные скрытые (латентные) "
     "признаки позволяют оценить предпочтение для тех сочетаний, которые "
     "в истории отсутствуют."),
    ("Холодный старт",
     " – ситуация, при которой рекомендательная модель не может сформировать "
     "персональную выдачу из-за отсутствия истории взаимодействий: "
     "для нового пользователя обученная модель не располагает сведениями "
     "и возвращает одинаковую оценку для всех объектов."),
    ("Кластеризация",
     " – разделение множества объектов на группы (кластеры) по сходству их "
     "признаков без заранее заданной разметки. В работе применяется метод "
     "k-means, относящий гостя к сегменту по среднему чеку и распределению "
     "его заказов по категориям меню."),
]


def style(p):
    pf = p.paragraph_format
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing = 1.5
    pf.space_before = pf.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def run(p, text, bold=False):
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold = FONT, SIZE, bold
    # для кириллицы Word смотрит не только на ascii-шрифт
    rpr = r._element.get_or_add_rPr().rFonts
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rpr.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + attr,
            FONT,
        )
    return r


def main():
    doc = Document(OUTPUT)
    paras = doc.paragraphs

    start = next(i for i, p in enumerate(paras)
                 if p.text.strip().upper() == "ТЕРМИНЫ И ОПРЕДЕЛЕНИЯ")
    end = next(i for i, p in enumerate(paras)
               if i > start and p.style.name.startswith("Heading"))

    old = [p for p in paras[start + 1:end] if p.text.strip()]
    for p in old:
        p._element.getparent().remove(p._element)

    anchor = paras[start]._element
    for term, body in TERMS:
        p = doc.add_paragraph()
        anchor.addnext(p._element)
        anchor = p._element
        style(p)
        run(p, term)
        run(p, body)

    doc.save(OUTPUT)
    print(f"Терминов удалено: {len(old)}, добавлено: {len(TERMS)}")
    print(f"Готово: {OUTPUT}")


if __name__ == "__main__":
    main()
